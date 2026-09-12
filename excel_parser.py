"""Read-only, label/merge based parser. No AI, no workbook writes or DEMO data."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import openpyxl
from openpyxl.cell.cell import MergedCell

PARSER_VERSION = "2.0.1"
DEFAULT_FOLDER = Path(__file__).parent / "訓練班別計畫表"


def normalized(value):
    return re.sub(r"\s+", "", str(value if value is not None else "")).replace("&nbsp", "")


def raw_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError("布林值不是數值")
    try:
        result = Decimal(str(value).strip().replace(",", ""))
        if not result.is_finite():
            raise ValueError("數值不是有限值")
        return result
    except InvalidOperation as exc:
        raise ValueError("非可辨識數值") from exc


def parse_date(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    match = re.fullmatch(r"[自至]?\s*(\d{3,4})[/年-](\d{1,2})[/月-](\d{1,2})日?", str(value).strip())
    if not match:
        raise ValueError("無法辨識日期")
    y, m, d = map(int, match.groups())
    return date(y + 1911 if y < 1911 else y, m, d).isoformat()


class FormatError(ValueError):
    pass


class SheetReader:
    def __init__(self, sheet, data):
        self.sheet, self.data = sheet, data
        self.anchors = [c for row in sheet for c in row if not isinstance(c, MergedCell)]
        self.merges = {}
        for area in sheet.merged_cells.ranges:
            for row in sheet.iter_rows(min_row=area.min_row, max_row=area.max_row,
                                       min_col=area.min_col, max_col=area.max_col):
                for c in row:
                    self.merges[c.coordinate] = area

    def find(self, text, prefix=False, required=True, min_row=1, max_row=None, max_col=None):
        target = normalized(text)
        found = [c for c in self.anchors if c.row >= min_row
                 and (max_row is None or c.row <= max_row)
                 and (max_col is None or c.column <= max_col)
                 and (normalized(c.value).startswith(target) if prefix else normalized(c.value) == target)]
        if len(found) != 1:
            if not required and not found:
                return None
            raise FormatError(f"表頭「{text}」{'找不到' if not found else '不唯一'}")
        return found[0]

    def area(self, cell):
        return self.merges.get(cell.coordinate)

    def right(self, label):
        area = self.area(label)
        return self.sheet.cell(label.row, (area.max_col if area else label.column) + 1)

    def capture(self, key, cell, convert=None):
        area = self.area(cell)
        if isinstance(cell, MergedCell) and area:
            cell = self.sheet.cell(area.min_row, area.min_col)
        raw = cell.value
        status = "原表空白" if raw is None or str(raw).strip() == "" else "已取得"
        value = None if status == "原表空白" else raw_value(raw)
        try:
            if cell.data_type in ("f", "e"):
                raise ValueError("公式或 Excel 錯誤值尚不支援；不採用可能過期的快取")
            if value is not None and convert:
                value = convert(raw)
        except (ValueError, TypeError) as exc:
            value, status = None, "無法解析"
            self.data["parse_report"]["errors"].append({"field": key, "message": str(exc)})
        self.data["field_sources"][key] = [{
            "file": self.data["source"]["relative_path"], "sheet": self.sheet.title,
            "cell": cell.coordinate, "range": str(area) if area else cell.coordinate,
            "raw_value": raw_value(raw), "status": status,
        }]
        return value

    def labeled(self, key, label, convert=None):
        return self.capture(key, self.right(self.find(label)), convert)

    def block(self, key, start, end, first_col=4, last_col=18):
        entries, refs = [], []
        for c in self.anchors:
            if start <= c.row <= end and first_col <= c.column <= last_col and c.value is not None:
                subkey = key + "." + c.coordinate
                value = self.capture(subkey, c)
                refs.extend(self.data["field_sources"].pop(subkey))
                if value is not None:
                    entries.append(str(value))
        self.data["field_sources"][key] = refs or [{
            "file": self.data["source"]["relative_path"], "sheet": self.sheet.title,
            "range": f"{openpyxl.utils.get_column_letter(first_col)}{start}:{openpyxl.utils.get_column_letter(last_col)}{end}",
            "raw_value": None, "status": "原表空白"}]
        return "\n".join(entries) or None


def parse_excel_bytes(content, filename, relative_path=None):
    digest = sha256(content).hexdigest()
    data = {"schema_version": "2.0", "case_id": None,
            "source": {"relative_path": relative_path or filename, "filename": filename,
                       "file_sha256": digest, "parser_version": PARSER_VERSION, "sheet": None},
            **{k: {} for k in ("plan", "organization", "course", "objectives", "audience",
                              "facilities", "budget", "other_sections", "derived", "field_sources")},
            "sessions": [], "instructors": [], "assistants": [],
            "parse_report": {"status": "failed", "errors": [], "warnings": []}}
    workbook = None
    try:
        workbook = openpyxl.load_workbook(BytesIO(content), data_only=False)
        candidates = [s for s in workbook if any(normalized(c.value) == "課程名稱"
                      for row in s for c in row if not isinstance(c, MergedCell))]
        if len(candidates) != 1:
            raise FormatError("必須有且只有一張可辨識的課程計畫表；不自動選取不明工作表")
        s = candidates[0]
        data["source"].update(sheet=s.title, merged_ranges=[str(a) for a in s.merged_cells.ranges])
        identity = f"{relative_path or filename}|{digest}|{s.title}|{PARSER_VERSION}"
        data["case_id"] = "EXCEL-" + sha256(identity.encode()).hexdigest()[:24]
        r = SheetReader(s, data)
        # Basic fields are located by label, not by the audited row numbers.
        fields = {"organization.name": "訓練單位名稱", "course.name": "課程名稱",
                  "course.occupation": "訓練職類", "course.competency": "訓練職能",
                  "course.industry": "業別分類代碼", "course.delivery": "辦理方式",
                  "course.policy_industry": "政策性產業", "course.time_description": "上課時間",
                  "course.enrollment": "訓練人數", "course.declared_hours": "訓練時數",
                  "course.weeks": "訓練週數"}
        for key, label in fields.items():
            group, name = key.split(".")
            data[group][name] = r.labeled(key, label, number if name in ("enrollment", "declared_hours", "weeks") else None)
        period = r.find("起迄日期")
        start = r.right(period)
        data["course"]["start_date"] = r.capture("course.start_date", start, parse_date)
        data["course"]["end_date"] = r.capture("course.end_date", s.cell(start.row + 1, start.column), parse_date)
        top_end = r.find("訓練單位名稱").row - 1
        top = [c for c in r.anchors if c.row <= top_end and c.value is not None]
        plan_cell = next((c for c in top if "產業人才投資計畫" in str(c.value) or "提升勞工自主學習計畫" in str(c.value)), None)
        year_cell = next((c for c in top if "年度" in str(c.value) and "訓練班別計畫表" in str(c.value)), None)
        if plan_cell is None or year_cell is None:
            raise FormatError("無法辨識計畫別／年度表頭")
        data["plan"]["name"] = r.capture("plan.name", plan_cell)
        year_raw = r.capture("plan.period_raw", year_cell)
        match = re.search(r"(\d+)年度\s*(上半年|下半年)", year_raw)
        data["plan"].update(year_roc=int(match[1]) if match else None, term=match[2] if match else None)
        data["field_sources"]["plan.year_roc"] = data["field_sources"]["plan.period_raw"]
        data["field_sources"]["plan.term"] = data["field_sources"]["plan.period_raw"]
        if not match:
            data["parse_report"]["errors"].append({"field": "plan.period_raw", "message": "年度期別無法拆分"})
        branch = next((c for c in top if "分署" in str(c.value)), None)
        data["plan"]["branch"] = r.capture("plan.branch", branch) if branch else None
        needs = r.find("訓練需求調查", prefix=True)
        objectives = r.find("訓練目標", prefix=True)
        teachers = r.find("師資(", prefix=True)
        assistants = r.find("助教(", prefix=True)
        audience = r.find("學員資格(", prefix=True)
        header = r.find("課程進度/內容")
        evaluation = r.find("訓練績效評估")
        fees = r.find("訓練費")
        if not needs.row < objectives.row < teachers.row < assistants.row < audience.row < header.row < evaluation.row < fees.row:
            raise FormatError("區塊順序不在目前支援格式內")
        data["objectives"]["needs_raw"] = r.block("objectives.needs_raw", needs.row, objectives.row - 1)
        data["objectives"]["goals_raw"] = r.block("objectives.goals_raw", objectives.row, teachers.row - 1)
        for group, begin, end in (("instructors", teachers.row, assistants.row), ("assistants", assistants.row, audience.row)):
            columns = {k: r.find(label, min_row=begin, max_row=begin).column
                       for k, label in {"name": "姓名", "education": "學歷", "experience": "工作經驗與年資",
                                        "certificates": "相關證照", "specialty": "專業領域"}.items()}
            columns["qualification"] = r.find("授課師資條件" if group == "instructors" else "助教條件", min_row=begin, max_row=begin).column
            r.block(group, begin + 1, end - 1)
            for row in range(begin + 1, end):
                cells = [s.cell(row, col) for col in columns.values()]
                if not any(c.value is not None and normalized(c.value) != "(依計畫師資及助教資格標準表)" for c in cells):
                    continue
                record = {k: r.capture(f"{group}.{len(data[group])}.{k}", s.cell(row, col)) for k, col in columns.items()}
                data[group].append(record)
        for offset, key in enumerate(("education_raw", "age_raw", "qualification_raw")):
            data["audience"][key] = r.capture("audience." + key, s.cell(audience.row + offset, r.right(audience).column))
        facility = r.find("裝備與設施")
        teaching = r.find("教學方法")
        data["facilities"]["raw"] = r.block("facilities.raw", facility.row, teaching.row - 1, 2)
        data["other_sections"]["teaching_raw"] = r.block("other_sections.teaching_raw", teaching.row, header.row - 1)
        data["other_sections"]["evaluation_raw"] = r.block("other_sections.evaluation_raw", evaluation.row, fees.row - 1, 2)
        data["other_sections"]["after_budget_raw"] = r.block("other_sections.after_budget_raw", fees.row + 3, s.max_row, 1)
        # Column headers identify the skill-hours variant; no file-number dispatch.
        skill = r.find("技檢訓練(時數)", required=False, min_row=header.row, max_row=header.row)
        data["source"]["format_variant"] = "skill_hours" if skill else "standard"
        cols = {k: r.find(label, min_row=header.row, max_row=header.row).column for k, label in {
            "date": "日期", "hours": "時數", "content": "課程進度/內容", "kind": "學/術科",
            "location": "授課地點", "teacher": "授課教師", "assistant": "助教", "remote": "遠距", "outdoor": "室外"}.items()}
        time_col = r.find("授課時間", min_row=header.row, max_row=header.row).column
        hours_col = cols["hours"]
        # Date/hour merge anchors define records, including records with blank dates/hours.
        data_start = max((r.area(s.cell(header.row, col)).max_row if r.area(s.cell(header.row, col)) else header.row) for col in (cols["date"], hours_col)) + 1
        row = data_start
        while row < evaluation.row:
            date_cell, hour_cell = s.cell(row, cols["date"]), s.cell(row, hours_col)
            areas = [r.area(date_cell), r.area(hour_cell)]
            end = max([row] + [a.max_row for a in areas if a])
            if end >= evaluation.row:
                raise FormatError("課表合併範圍跨越區塊邊界")
            index = len(data["sessions"])
            record = {k: r.capture(f"sessions.{index}.{k}", s.cell(row, col), parse_date if k == "date" else number if k == "hours" else None) for k, col in cols.items()}
            record["weekday"] = r.capture(f"sessions.{index}.weekday", s.cell(row, time_col))
            record["time_raw"] = r.capture(f"sessions.{index}.time_raw", s.cell(row, time_col + 1))
            record["skill_hours"] = r.capture(f"sessions.{index}.skill_hours", s.cell(row, skill.column), number) if skill else None
            data["sessions"].append(record)
            row = end + 1
        if not data["sessions"]:
            data["parse_report"]["warnings"].append("課表未提供資料")
        for key, label in {"subsidy": "政府補助", "self_paid": "學員自付", "total": "總計"}.items():
            cell = r.find(label, min_row=fees.row, max_row=fees.row + 2)
            class_cell = r.right(cell)
            class_unit = r.right(class_cell)
            person_cell = r.right(class_unit)
            for basis, c in (("per_class", class_cell), ("per_person", person_cell)):
                data["budget"].setdefault(basis, {})[key] = r.capture(f"budget.{basis}.{key}", c, number)
        for key, label in {"fixed_total": "固定費用總額", "materials_total": "材料費用總額", "hourly_cost": "固定費用單一人時成本", "materials_ratio_raw": "材料費占比"}.items():
            cell = r.find(label, prefix=True, required=False, min_row=audience.row, max_row=header.row - 1)
            if cell:
                def cost(v):
                    value = re.split(r"[:：]", str(v), maxsplit=1)[-1].strip().removesuffix("元").strip()
                    return number(value)
                data["budget"][key] = r.capture("budget." + key, cell, None if key.endswith("raw") else cost)
            else:
                data["budget"][key] = None
                data["field_sources"]["budget." + key] = [{"file": data["source"]["relative_path"], "sheet": s.title, "range": None, "raw_value": None, "status": "找不到欄位"}]
        cost_label = r.find("訓練費用編列說明")
        data["budget"]["details_raw"] = r.block("budget.details_raw", audience.row, header.row - 1, r.right(cost_label).column)
        hours = [x["hours"] for x in data["sessions"]]
        total = sum(hours, Decimal(0)) if hours and all(x is not None for x in hours) else None
        declared = data["course"]["declared_hours"]
        data["derived"] = {"session_count": len(hours), "schedule_hours": total,
                           "hours_difference": total - declared if total is not None and declared is not None else None,
                           "schedule_hours_inputs": [f"sessions.{i}.hours" for i in range(len(hours))],
                           "method": "各筆課表時數相加；任一時數缺漏時不產生完整合計"}
        for group in ("instructors", "assistants"):
            if not data[group]:
                data["parse_report"]["warnings"].append(f"{group}：原表未提供實際人員資料")
        data["parse_report"]["status"] = "partial" if data["parse_report"]["errors"] else "success"
    except Exception as exc:
        data["parse_report"]["status"] = "failed"
        data["parse_report"]["errors"].append({"field": "workbook", "message": f"{type(exc).__name__}: {exc}"})
    finally:
        if workbook:
            workbook.close()
    return data


def scan_folder(folder=DEFAULT_FOLDER):
    """Catalogue only; no review state is initialized."""
    folder = Path(folder)
    if not folder.is_dir():
        raise FileNotFoundError(f"找不到資料夾：{folder}")
    paths = [p for p in folder.rglob("*") if p.is_file() and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")]
    return sorted(paths, key=lambda p: (int(re.match(r"\((\d+)\)", p.name)[1]) if re.match(r"\((\d+)\)", p.name) else 9999, p.name))


def parse_excel(path, folder=DEFAULT_FOLDER):
    path = Path(path)
    relative = str(path.resolve().relative_to(Path(folder).resolve()))
    try:
        content = path.read_bytes()
    except OSError as exc:
        result = parse_excel_bytes(b"", path.name, relative)
        result["parse_report"]["errors"] = [{"field": "workbook", "message": str(exc)}]
        return result
    return parse_excel_bytes(content, path.name, relative)
